import base64
import io
import json
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
import zipfile
from pathlib import Path
from http.server import ThreadingHTTPServer
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from agent.importers import extract_file
from agent.analysis import local_candidates
from agent.domain import sample_project
from agent.exports import make_vsdx
from agent.server import create_handler
from agent.store import Store


def package(parts):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, text in parts.items():
            archive.writestr(name, text)
    return buf.getvalue()


def presentation_parts():
    # Filename order deliberately differs from presentation order.
    return {
        'ppt/presentation.xml': '<p:presentation xmlns:p="urn:p" xmlns:r="urn:r"><p:sldIdLst><p:sldId r:id="second"/><p:sldId r:id="first"/></p:sldIdLst></p:presentation>',
        'ppt/_rels/presentation.xml.rels': '<Relationships><Relationship Id="second" Type="urn:x/slide" Target="slides/slide2.xml"/><Relationship Id="first" Type="urn:x/slide" Target="slides/slide1.xml"/></Relationships>',
        'ppt/slides/slide2.xml': '<p:sld xmlns:p="urn:p" xmlns:a="urn:a" show="0"><p:sp><p:cNvPr id="7"/><a:p><a:r><a:t>VLAN</a:t></a:r><a:r><a:t>20を分離</a:t></a:r></a:p></p:sp><a:tbl><a:tr><a:tc><a:p><a:t>REQ-01</a:t></a:p></a:tc><a:tc><a:p><a:t>認証を二重化</a:t></a:p></a:tc></a:tr></a:tbl></p:sld>',
        'ppt/slides/slide1.xml': '<sld><sp><cNvPr id="1"/><p><t>最後のスライド</t></p></sp></sld>',
        'ppt/slides/_rels/slide2.xml.rels': '<Relationships><Relationship Id="notes" Type="urn:x/notesSlide" Target="../notesSlides/notesSlide1.xml"/><Relationship Id="outside" TargetMode="External" Type="urn:x/notesSlide" Target="https://example.invalid/secret"/></Relationships>',
        'ppt/notesSlides/notesSlide1.xml': '<notes><sp><ph type="sldNum"/><p><t>999</t></p></sp><sp><ph type="body"/><p><t>認証の補足</t></p></sp></notes>',
    }


class ImporterTests(unittest.TestCase):
    def test_pptx_order_runs_tables_notes_hidden_and_no_network(self):
        with patch('urllib.request.urlopen', side_effect=AssertionError('network not allowed')):
            text = extract_file('input.PPTX', package(presentation_parts()))
        self.assertIn('スライド1（非表示） 図形7: VLAN20を分離', text)
        self.assertIn('表1 行1: REQ-01 | 認証を二重化', text)
        self.assertIn('スライド1（非表示） ノート: 認証の補足', text)
        self.assertIn('スライド2 図形1: 最後のスライド', text)
        self.assertNotIn('999', text)
        result = local_candidates([dict(name='input.pptx', body=text)])
        self.assertIn('スライド1', result['requirements'][0]['source'])

    def test_external_or_missing_required_slide_is_rejected(self):
        parts = presentation_parts()
        parts['ppt/_rels/presentation.xml.rels'] = '<Relationships/>'
        with self.assertRaisesRegex(ValueError, '資料を読み取れません'):
            extract_file('x.pptx', package(parts))

    def test_no_image_only_success(self):
        parts = presentation_parts()
        parts['ppt/slides/slide2.xml'] = '<sld><pic/></sld>'
        parts['ppt/slides/slide1.xml'] = '<sld><pic/></sld>'
        del parts['ppt/slides/_rels/slide2.xml.rels']
        with self.assertRaisesRegex(ValueError, '文字を抽出できません'):
            extract_file('image.pptx', package(parts))

    def test_vsdx_roundtrip_registered_connections(self):
        with tempfile.TemporaryDirectory() as directory:
            file = Path(directory) / 'topology.vsdx'
            make_vsdx(file, sample_project()['design'])
            text = extract_file(file.name, file.read_bytes())
        self.assertIn('ページ1(L2構成図)', text)
        self.assertIn('DIST-SW01', text)
        self.assertIn('192.0.2.11/24', text)
        self.assertIn('.BeginX → 図形1.PinX', text)

    def test_visio_nested_shapes_properties_master_and_vdx(self):
        page = '<Page Name="現行"><Shapes><Shape ID="1" NameU="Group"><Shapes><Shape ID="2" Master="5"><Text>BR1<cp IX="0"/>-SW01</Text><Section N="Property"><Row N="IP"><Cell N="Label" V="管理IP"/><Cell N="Value" V="192.0.2.20"/></Row></Section></Shape></Shapes></Shape></Shapes></Page>'
        text = extract_file('x.vdx', ('<VisioDocument><Pages>' + page + '</Pages></VisioDocument>').encode())
        self.assertIn('図形2: BR1-SW01', text)
        self.assertIn('属性管理IP: 192.0.2.20', text)
        self.assertIn('Master=5', text)
        self.assertIn('未展開', text)

    def test_xlsx_formulas_not_silently_dropped(self):
        book = Workbook()
        book.active.title = '台帳'
        book.active['A1'] = '管理IP'
        book.active['B1'] = '192.0.2.1'
        book.active['A2'] = '=1+2'
        hidden = book.create_sheet('旧情報')
        hidden.sheet_state = 'hidden'
        hidden['C3'] = '旧VLAN'
        buf = io.BytesIO(); book.save(buf); book.close()
        text = extract_file('x.xlsx', buf.getvalue())
        self.assertIn('台帳: A1=管理IP | B1=192.0.2.1', text)
        self.assertIn('A2==1+2 [数式・保存済み値なし／未再計算]', text)
        self.assertIn('旧情報（非表示）: C3=旧VLAN', text)

    def test_excel_sparse_huge_dimensions_rejected(self):
        book = Workbook(); book.active['XFD20000'] = 'sparse'
        buf = io.BytesIO(); book.save(buf); book.close()
        with self.assertRaisesRegex(ValueError, '256列'):
            extract_file('x.xlsx', buf.getvalue())

    def test_docx_and_text_encodings(self):
        book = Document(); book.add_paragraph('業務とゲストを分離')
        book.add_table(rows=1, cols=1).cell(0, 0).text = '管理VLAN99'
        buf = io.BytesIO(); book.save(buf)
        self.assertIn('表1 行1: 管理VLAN99', extract_file('x.docx', buf.getvalue()))
        for encoding in ('utf-8-sig', 'utf-16', 'cp932'):
            self.assertIn('L1: 認証を二重化', extract_file('x.txt', '認証を二重化'.encode(encoding)))

    def test_legacy_and_corrupt_formats_actionable(self):
        for ext in ('.doc', '.xls', '.ppt', '.vsd'):
            with self.subTest(ext=ext), self.assertRaisesRegex(ValueError, '名前を付けて保存'):
                extract_file('x' + ext, b'legacy binary')
        for ext in ('.docx', '.xlsx', '.pptx', '.vsdx'):
            with self.subTest(ext=ext), self.assertRaisesRegex(ValueError, '破損'):
                extract_file('x' + ext, b'not a zip')

    def test_xml_entities_traversal_and_resource_limits(self):
        dtd = '<!DOCTYPE x [<!ENTITY x "secret">]><VisioDocument>&x;</VisioDocument>'
        for encoding in ('utf-8', 'utf-16'):
            with self.assertRaisesRegex(ValueError, 'DTD'):
                extract_file('x.vdx', dtd.encode(encoding))
        parts = presentation_parts()
        parts['ppt/_rels/presentation.xml.rels'] = parts['ppt/_rels/presentation.xml.rels'].replace('slides/slide2.xml', '../../outside.xml')
        with self.assertRaisesRegex(ValueError, 'パッケージ外'):
            extract_file('x.pptx', package(parts))
        with self.assertRaisesRegex(ValueError, '展開サイズ'):
            extract_file('bomb.pptx', package({'x': b'0' * (40 * 1024 * 1024 + 1)}))
        with self.assertRaisesRegex(ValueError, '16万字'):
            extract_file('x.txt', ('あ' * 160001).encode())


class ImportApiTests(unittest.TestCase):
    def test_upload_formats_and_rejection_does_not_add_source(self):
        with tempfile.TemporaryDirectory() as directory:
            store = Store(Path(directory)); hosts = set()
            server = ThreadingHTTPServer(('127.0.0.1', 0), create_handler(store, hosts))
            hosts.add(f'127.0.0.1:{server.server_port}')
            thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
            def request(endpoint, body=None):
                req = urllib.request.Request(f'http://127.0.0.1:{server.server_port}' + endpoint,
                    data=json.dumps(body).encode() if body is not None else None,
                    headers={'Content-Type': 'application/json', 'X-Workbench-Request': '1'})
                with urllib.request.urlopen(req) as response:
                    return json.loads(response.read())
            try:
                self.assertIn('.vsdx', request('/api/health')['input_formats'])
                project = request('/api/projects', {'name': '取込試験'})
                endpoint = '/api/projects/' + project['id'] + '/sources'
                request(endpoint, {'name': 'input.pptx', 'base64': base64.b64encode(package(presentation_parts())).decode()})
                with self.assertRaises(urllib.error.HTTPError) as caught:
                    request(endpoint, {'name': 'broken.vsdx', 'base64': base64.b64encode(b'bad').decode()})
                self.assertEqual(caught.exception.code, 400)
                sources = request(endpoint)
                self.assertEqual(len(sources), 1)
                self.assertIn('認証を二重化', sources[0]['body'])
            finally:
                server.shutdown(); server.server_close(); thread.join()
