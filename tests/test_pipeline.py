import copy
import io
import json
import os
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.request
import zipfile
from pathlib import Path
from unittest.mock import patch
from http.server import ThreadingHTTPServer

from docx import Document
from openpyxl import load_workbook
from agent.domain import Project
from agent.design_report import local_draft
from agent.pipeline import check_architecture, check_delivery, generate_pipeline
from agent.server import create_handler
from agent.store import Store


def fixture():
    requirements = [dict(id='REQ-1', text='2拠点の管理LANを分離する。', status='confirmed', source='要件.txt')]
    snapshot = dict(revision=1, approved_revision=None, project=Project(name='全工程検証', requirements=requirements).model_dump())
    sources = [dict(id='s1', name='要件.txt', body='2拠点の管理LANを分離する。')]
    trace = dict(requirement_ids=['REQ-1'], evidence=[dict(source_id='s1', quote=sources[0]['body'])], state='proposal', rationale='レビュー用の提案')
    a = dict(overview='2拠点を分離する案', sites=[], networks=[], nodes=[], parameters=[], connections=[], questions=['機種とOSを決定する'])
    for i in (1, 2):
        a['sites'].append(dict(trace, id=f'S{i}', name=f'拠点{i}', kind='branch'))
        a['networks'].append(dict(trace, id=f'N{i}', site_id=f'S{i}', name='管理', vrf='default', vlan=99, cidr=f'192.0.{i}.0/24', gateway=f'192.0.{i}.1'))
        a['nodes'].append(dict(trace, id=f'D{i}', hostname=f'BR{i}-SW01', site_id=f'S{i}', role='スイッチ', model=None, os_version=None, network_id=f'N{i}', management_ip=f'192.0.{i}.11'))
    a['parameters'] = [dict(trace, id='P1', node_id='D1', name='注記', value='=HYPERLINK("https://example.invalid")')]
    a['connections'] = [dict(trace, id='L1', a='D1', b='D2', a_port=None, b_port=None, kind='logical', purpose='閉域接続の検討')]
    step = dict(id='STEP-1', targets=['D1'], requirement_ids=['REQ-1'], action='設定を台帳と照合', expected='アドレスが一致', rollback='旧設定へ復元', owner='NW担当')
    d = dict(test_policy='検証環境で確認', tests=[dict(id='T1', requirement_ids=['REQ-1'], targets=['D1','MEASURE'], kind='normal', preconditions='検証機を使用', procedure='管理LANから疎通を確認', expected='許可された宛先のみ応答', evidence_to_collect='疎通ログ')],
        test_nodes=[dict(trace, id='MEASURE', hostname='TEST-PC', site_id='S1', role='測定端末', model=None, os_version=None, network_id=None, management_ip=None)],
        test_connections=[dict(trace,id='TL1',a='D1',b='MEASURE',a_port=None,b_port=None,kind='physical',purpose='測定用接続')],
        migration=[copy.deepcopy(step)], construction=[copy.deepcopy(step)], operations=[copy.deepcopy(step)],
        configs=[dict(node_id=f'D{i}',status='deferred',text='',reason='機種・OS未指定') for i in (1,2)],
        coverage=[dict(requirement_id='REQ-1',disposition='designed',design_ids=['N1','N2'],test_ids=['T1'],reason='拠点別サブネットと疎通試験')])
    review = dict(summary='機種未決のためレビューが必要', issues=[dict(severity='blocking',target='Config',finding='機種未決',required_action='機種とOSを確定')])
    return snapshot, sources, a, d, review


class PipelineTests(unittest.TestCase):
    def test_multisite_vlan_and_address_validation(self):
        s, sources, a, d, _ = fixture()
        check_delivery(d, a, s, sources)  # VLAN99 is valid in two different sites.
        for kind in ('overlap','gateway','host','source','reference','parameter','port'):
            b=copy.deepcopy(a)
            if kind=='overlap': b['networks'][1]['cidr']=b['networks'][0]['cidr']
            if kind=='gateway': b['networks'][0]['gateway']='203.0.113.1'
            if kind=='host': b['nodes'][1]['hostname']=b['nodes'][0]['hostname'].lower()
            if kind=='source': b['nodes'][0]['evidence'][0]['quote']='原文にない引用'
            if kind=='reference': b['nodes'][0]['network_id']='other'
            if kind=='parameter': b['parameters'].append(dict(b['parameters'][0],id='P2'))
            if kind=='port':
                b['connections'][0].update(kind='physical',a_port='Gi1',b_port='Gi1')
                b['connections'].append(dict(b['connections'][0],id='L2'))
            with self.subTest(kind=kind), self.assertRaises(ValueError): check_architecture(b,s,sources)

    def test_coverage_and_config_fail_closed(self):
        s, sources, a, d, _ = fixture()
        for kind in ('missing','test','unknown','duplicate','config','deferred','target','test_ip'):
            bad=copy.deepcopy(d)
            if kind=='missing': bad['coverage']=[]
            if kind=='test': bad['coverage'][0]['test_ids']=[]
            if kind=='unknown': bad['coverage'][0]['design_ids']=['OTHER']
            if kind=='duplicate': bad['coverage']*=2
            if kind=='config': bad['configs'][0].update(status='draft',text='hostname BR1-SW01')
            if kind=='deferred': bad['configs'][0]['text']='unexpected command'
            if kind=='target': bad['tests'][0]['targets']=['NO-NODE']
            if kind=='test_ip': bad['test_nodes'][0]['management_ip']='192.0.1.11'
            with self.subTest(kind=kind), self.assertRaises(ValueError): check_delivery(bad,a,s,sources)

    def test_full_pipeline_exports_shared_values_and_keeps_input(self):
        s, sources, a, d, review=fixture()
        original=copy.deepcopy(s); calls=[]; progress=[]
        replies=iter([a,local_draft(s,sources),d,review])
        def chat(prompt, context):
            calls.append((prompt,context))
            return json.dumps(next(replies),ensure_ascii=False)
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)/'job'
            result=generate_pipeline(s,sources,folder,'m365',chat,progress.append)
            self.assertEqual(len(calls),4)
            self.assertIn('共通設計案',calls[1][1])
            self.assertEqual(len(progress),5)
            self.assertEqual(result['blocking'],1)
            self.assertEqual(result['status'],'review_required')
            self.assertEqual(result['config_deferred'],2)
            self.assertEqual(s,original)
            with zipfile.ZipFile(folder.with_suffix('.zip')) as z:
                self.assertEqual(z.testzip(),None)
                self.assertIn('11_運用引継ぎ資料.docx',z.namelist())
                self.assertFalse(any(n.endswith('.cfg') for n in z.namelist()))
            wb=load_workbook(folder/'03_パラメータシート.xlsx')
            self.assertEqual(wb['機器台帳']['B2'].value,'BR1-SW01')
            self.assertEqual(wb['機器台帳']['F2'].value,'192.0.1.11')
            self.assertEqual(wb['全パラメータ']['D2'].data_type,'s')
            self.assertTrue(wb['全パラメータ']['D2'].value.startswith('='))
            wb.close()
            self.assertIn('192.0.1.11','\n'.join(p.text for p in Document(folder/'02b_詳細設計書.docx').paragraphs))
            self.assertIn('TEST-PC',(folder/'07_試験構成図.svg').read_text(encoding='utf-8'))
            self.assertNotIn('TEST-PC',(folder/'04_構成図.svg').read_text(encoding='utf-8'))
            self.assertIn('未実施',load_workbook(folder/'08_試験項目表.xlsx').active['I2'].value)

    def test_failure_does_not_publish_partial_zip(self):
        s,sources,a,d,review=fixture()
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)/'job'
            replies=iter([a,local_draft(s,sources),d,review])
            with patch('agent.pipeline.write_xlsx',side_effect=ValueError('writer failed')):
                with self.assertRaisesRegex(ValueError,'writer failed'):
                    generate_pipeline(s,sources,folder,'m365',lambda *_:json.dumps(next(replies)))
            self.assertEqual(list(Path(tmp).iterdir()),[])

    def test_no_unconfigured_or_local_fallback(self):
        s,sources,*_=fixture()
        with tempfile.TemporaryDirectory() as tmp,patch.dict(os.environ,{'LLM_MODEL':''}):
            for mode in ('local','ai','m365'):
                with self.subTest(mode=mode),self.assertRaises(ValueError):generate_pipeline(s,sources,Path(tmp)/'job',mode)

    def test_config_requires_model_and_os_in_referenced_input(self):
        s,sources,a,d,_=fixture()
        a['nodes'][0].update(model='C9200L',os_version='IOS XE 17.9.4',state='source')
        d['configs'][0].update(status='draft',text='hostname BR1-SW01')
        with self.assertRaisesRegex(ValueError,'機種・OSを関連する原文'):
            check_delivery(d,a,s,sources)
        s['project']['requirements'][0]['text']+=' C9200L / IOS XE 17.9.4を使用。'
        check_delivery(d,a,s,sources)
        d['configs'][0]['text']+='\x00'
        with self.assertRaisesRegex(ValueError,'制御文字'):check_delivery(d,a,s,sources)

    def test_http_pipeline_auth_revision_download_and_no_mutation(self):
        s,sources,a,d,review=fixture()
        with tempfile.TemporaryDirectory() as tmp:
            store=Store(Path(tmp)); item=store.create(s['project'])
            sid=store.add_source(item['id'],'要件.txt',sources[0]['body'])
            # Let the mock use actual source IDs from the API snapshot.
            for collection in (a['sites'],a['networks'],a['nodes'],a['parameters'],a['connections'],d['test_nodes'],d['test_connections']):
                for x in collection: x['evidence'][0]['source_id']=sid
            sources[0]['id']=sid
            replies=iter([a,local_draft(item,sources),d,review])
            hosts=set(); server=ThreadingHTTPServer(('127.0.0.1',0),create_handler(store,hosts))
            hosts.add(f'127.0.0.1:{server.server_port}'); base=f'http://127.0.0.1:{server.server_port}'
            threading.Thread(target=server.serve_forever,daemon=True).start()
            def call(path,body=None,headers=None):
                h={'Content-Type':'application/json','X-Workbench-Request':'1'};h.update(headers or {})
                request=urllib.request.Request(base+path,data=json.dumps(body).encode() if body is not None else None,headers=h)
                with urllib.request.urlopen(request,timeout=20) as r:return r.read(),r.headers
            endpoint='/api/projects/'+item['id']+'/pipeline'
            try:
                with patch.dict(os.environ,{'LLM_MODEL':''}):
                    for body,h,code in [({}, {},400),({'revision':True},{},400),({'revision':2},{},409),({'revision':1,'mode':'m365'},{},401),({'revision':1,'mode':'local'},{},400),({'revision':1,'mode':'ai'},{},400),({'revision':1},{'Origin':'http://evil.invalid'},403)]:
                        with self.subTest(body=body),self.assertRaises(urllib.error.HTTPError) as e:call(endpoint,body,h)
                        self.assertEqual(e.exception.code,code)
                with patch('agent.server.llm_settings',return_value={'configured':True}),patch('agent.pipeline.design_chat_json',side_effect=lambda *_:next(replies)):
                    job_id=json.loads(call(endpoint,{'revision':1,'mode':'ai'})[0])['job_id']
                    for _ in range(150):
                        job=store.job(job_id)
                        if job['status'] not in ('running','queued'):break
                        time.sleep(.05)
                    self.assertEqual(job['status'],'completed',job)
                    data,headers=call(job['result']['download_url'])
                    self.assertEqual(headers['Content-Type'],'application/zip')
                    self.assertIn('design-model.json',zipfile.ZipFile(io.BytesIO(data)).namelist())
                    self.assertEqual(store.get(item['id']),item)
            finally:server.shutdown();server.server_close()


if __name__=='__main__':unittest.main()
