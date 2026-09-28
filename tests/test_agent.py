from __future__ import annotations

import copy
import io
import json
import os
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.request
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

from docx import Document
from openpyxl import load_workbook
from pydantic import ValidationError

from agent.analysis import ai_candidates, extract_file, llm_settings
from agent.domain import Project, config_text, sample_project, validate_project
from agent.exports import export_bundle
from agent.server import create_handler
from agent.store import Conflict, Store


class DomainTests(unittest.TestCase):
    def setUp(self): self.p=sample_project()
    def codes(self): return {x["code"] for x in validate_project(self.p)}

    def test_sample_is_statically_valid(self):
        self.assertFalse([x for x in validate_project(self.p) if x["level"] in ("error","warning")])

    def test_duplicate_ip(self):
        self.p["design"]["devices"][1]["management_ip"]="192.0.2.11/24"
        self.assertIn("IP_DUP",self.codes())

    def test_overlap_is_scoped_to_vrf(self):
        self.p["design"]["vlans"][1]["subnet"]="198.51.100.0/25"
        self.assertIn("SUBNET_OVERLAP",self.codes())
        self.p["design"]["vlans"][1]["vrf"]="guest"
        self.assertNotIn("SUBNET_OVERLAP",self.codes())

    def test_vlan_reference_and_link_mismatch(self):
        self.p["design"]["devices"][1]["ports"][0]["vlans"]=[400]
        self.assertTrue({"PORT_VLAN","LINK_MISMATCH"}.issubset(self.codes()))

    def test_duplicate_physical_endpoint(self):
        self.p["design"]["links"].append(copy.deepcopy(self.p["design"]["links"][0]))
        self.assertIn("LINK_DUP",self.codes())

    def test_invalid_gateway(self):
        self.p["design"]["devices"][0]["gateway"]="203.0.113.1"
        self.assertIn("GATEWAY",self.codes())

    def test_network_address_rejected(self):
        self.p["design"]["devices"][0]["management_ip"]="192.0.2.0/24"
        self.assertIn("IP_INVALID",self.codes())

    def test_unknown_adapter_is_blocked(self):
        self.p["design"]["devices"][0]["model"]="Unknown"
        self.assertIn("ADAPTER_UNSUPPORTED",self.codes())

    def test_config_injection_rejected(self):
        self.p["design"]["devices"][0]["ports"][0]["description"]="uplink\nreload"
        with self.assertRaises(ValidationError): Project.model_validate(self.p)

    def test_config_uses_exact_design_values(self):
        d=self.p["design"]; c=config_text(d["devices"][0],d)
        self.assertIn("ip address 192.0.2.11 255.255.255.0",c)
        self.assertIn("switchport trunk allowed vlan 10,20,99",c)
        self.assertNotIn("write memory",c)


class StoreTests(unittest.TestCase):
    def test_source_delete_is_scoped_and_preserves_other_sources(self):
        with tempfile.TemporaryDirectory() as tmp:
            store=Store(Path(tmp)); p=store.create(sample_project()); other=store.create(sample_project())
            a=store.add_source(p['id'],'same.txt','L1: 分離する')
            b=store.add_source(p['id'],'same.txt','L1: 分離する')
            with self.assertRaises(KeyError): store.delete_source(other['id'],a,1)
            p['project']['requirements']=[dict(id='A',text='分離する',source='same.txt / L1',source_id=a,quote='分離する',status='confirmed'),dict(id='B',text='分離する',source='same.txt / L1',source_id=b,quote='分離する',status='confirmed')]
            store.save(p['id'],1,p['project']);store.approve(p['id'],2)
            result=store.delete_source(p['id'],a,2)
            self.assertEqual(result['affected_requirements'],1)
            current=store.get(p['id'])
            self.assertEqual(current['revision'],3);self.assertIsNone(current['approved_revision'])
            self.assertEqual(current['project']['requirements'][0]['status'],'unresolved')
            self.assertEqual(current['project']['requirements'][0]['quote'],'分離する')
            self.assertIn('元資料削除',current['project']['requirements'][0]['source'])
            self.assertEqual(current['project']['requirements'][1]['status'],'confirmed')
            self.assertEqual([s['id'] for s in store.sources(p['id'])],[b])
            with self.assertRaises(KeyError): store.delete_source(p['id'],a,3)

    def test_source_delete_conflict_rolls_back_and_legacy_refs_are_flagged(self):
        with tempfile.TemporaryDirectory() as tmp:
            store=Store(Path(tmp)); p=store.create(sample_project())
            sid=store.add_source(p['id'],'legacy.txt','L1: 要件')
            p['project']['requirements']=[dict(id='R',text='要件',source='legacy.txt / L1',quote='要件',status='confirmed')]
            store.save(p['id'],1,p['project'])
            with self.assertRaises(Conflict):store.delete_source(p['id'],sid,1)
            self.assertEqual(len(store.sources(p['id'])),1)
            self.assertEqual(store.get(p['id'])['project']['requirements'][0]['status'],'confirmed')
            self.assertEqual(store.delete_source(p['id'],sid,2)['affected_requirements'],1)

    def test_revision_conflict_and_approval_invalidation(self):
        with tempfile.TemporaryDirectory() as tmp:
            store=Store(Path(tmp)); first=store.create(sample_project())
            store.approve(first["id"],1)
            saved=store.save(first["id"],1,first["project"])
            self.assertEqual(saved["revision"],2);self.assertIsNone(saved["approved_revision"])
            with self.assertRaises(Conflict): store.save(first["id"],1,first["project"])
            with store.connect() as db: self.assertEqual(db.execute("SELECT count(*) FROM revisions").fetchone()[0],2)


class ExtractionTests(unittest.TestCase):
    def test_text_source_locations(self):
        self.assertIn("L2: 管理VLAN",extract_file("sample.txt","要件\n管理VLAN".encode()))

    def test_docx_source_locations(self):
        doc=Document();doc.add_paragraph("VLANで分離する");buf=io.BytesIO();doc.save(buf)
        self.assertIn("段落1: VLANで分離する",extract_file("sample.docx",buf.getvalue()))

    def test_remote_ai_requires_explicit_opt_in(self):
        with patch.dict(os.environ,{"LLM_BASE_URL":"https://example.com/v1","LLM_MODEL":"test","ALLOW_REMOTE_LLM":"false"}):
            self.assertFalse(llm_settings()["configured"])

    def test_ai_evidence_verified_against_actual_source(self):
        sources=[dict(id="source1",name="rfp.txt",body="L1: VLANで分離する")]
        reply={"requirements":[{"text":"端末をVLANで分離する","quote":"VLANで分離する","source_id":"source1","category":"機能"}],"advice":"台数の確認が必要"}
        class Fake(BaseHTTPRequestHandler):
            def do_POST(self):
                self.rfile.read(int(self.headers["Content-Length"]))
                output=json.dumps({"choices":[{"message":{"content":json.dumps(reply,ensure_ascii=False)},"finish_reason":"stop"}]}).encode()
                self.send_response(200);self.send_header("Content-Length",str(len(output)));self.end_headers();self.wfile.write(output)
            def log_message(self,*a): pass
        server=ThreadingHTTPServer(("127.0.0.1",0),Fake)
        thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        try:
            with patch.dict(os.environ,{"LLM_BASE_URL":f"http://127.0.0.1:{server.server_port}/v1","LLM_MODEL":"test","LLM_API_KEY":""}):
                result=ai_candidates(sources);self.assertEqual(result["requirements"][0]["status"],"candidate")
                reply["requirements"][0]["quote"]="存在しない原文"
                with self.assertRaisesRegex(ValueError,"根拠"): ai_candidates(sources)
        finally:server.shutdown();server.server_close()


class ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory(); cls.store=Store(Path(cls.temp.name))
        hosts=set(); cls.server=ThreadingHTTPServer(("127.0.0.1",0),create_handler(cls.store,hosts))
        hosts.add(f"127.0.0.1:{cls.server.server_port}")
        cls.base=f"http://127.0.0.1:{cls.server.server_port}"
        threading.Thread(target=cls.server.serve_forever,daemon=True).start()
    @classmethod
    def tearDownClass(cls): cls.server.shutdown();cls.server.server_close();cls.temp.cleanup()

    def req(self,path,body=None,headers=None):
        h={"Content-Type":"application/json","X-Workbench-Request":"1"};h.update(headers or {})
        r=urllib.request.Request(self.base+path,data=json.dumps(body).encode() if body is not None else None,headers=h)
        with urllib.request.urlopen(r) as response: return json.loads(response.read())

    def await_job(self,jid):
        for _ in range(200):
            job=self.req("/api/jobs/"+jid)
            if job["status"] not in ("running","queued"): return job
            time.sleep(.05)
        self.fail("job timed out")

    def test_cross_origin_and_host_blocked(self):
        for headers in [{"Origin":"http://evil.example"},{"Host":"evil.example"},{"X-Workbench-Request":""}]:
            with self.assertRaises(urllib.error.HTTPError) as error:self.req("/api/projects",{"sample":True},headers)
            self.assertEqual(error.exception.code,403)

    def test_extraction_requires_explicit_apply(self):
        p=self.req("/api/projects",{"name":"テスト案件"});base="/api/projects/"+p["id"]
        self.req(base+"/sources",{"name":"rfp.txt","text":"業務端末とゲスト端末を分離する。"})
        job=self.await_job(self.req(base+"/analyze",{"mode":"local"})["job_id"])
        self.assertEqual(job["status"],"completed")
        self.assertEqual(self.req(base)["project"]["requirements"],[])
        applied=self.req("/api/jobs/"+job["id"]+"/apply",{})
        self.assertEqual(len(applied["project"]["requirements"]),1)
        with self.assertRaises(urllib.error.HTTPError) as error:self.req("/api/jobs/"+job["id"]+"/apply",{})
        self.assertEqual(error.exception.code,409)

    def test_exports_gate_config_and_preserve_values(self):
        p=self.req("/api/projects",{"sample":True});base="/api/projects/"+p["id"]
        job=self.await_job(self.req(base+"/generate",{"revision":1})["job_id"])
        self.assertEqual(job["status"],"completed",job)
        self.assertFalse(job["result"]["config_included"])
        self.req(base+"/approve",{"revision":1})
        job=self.await_job(self.req(base+"/generate",{"revision":1})["job_id"])
        self.assertTrue(job["result"]["config_included"])
        folder=self.store.root/"exports"/job["id"]
        with zipfile.ZipFile(folder.with_suffix(".zip")) as z:
            self.assertIn("05_Config/DIST-SW01.cfg",z.namelist())
            self.assertIn(b"192.0.2.11",z.read("05_Config/DIST-SW01.cfg"))
        wb=load_workbook(folder/"03_パラメータシート.xlsx",data_only=True)
        self.assertEqual(wb["アドレス台帳"]["C2"].value,"192.0.2.11/24")
        wb.close()
        for file in folder.glob("*.docx"):
            self.assertTrue(Document(file).paragraphs[0].text)
        for file in folder.glob("*.vsdx"):
            with zipfile.ZipFile(file) as z:
                for name in z.namelist(): ET.fromstring(z.read(name))
        self.assertIn("network-design-v1.zip",urllib.request.urlopen(self.base+job["result"]["download_url"]).headers["Content-Disposition"])

    def test_source_delete_api_invalidates_previous_analysis(self):
        project=self.req('/api/projects',{'name':'削除試験'})
        base='/api/projects/'+project['id']
        source=self.req(base+'/sources',{'name':'test.txt','text':'業務ネットワークを分離する。'})
        job=self.await_job(self.req(base+'/analyze',{'mode':'local'})['job_id'])
        endpoint=base+'/sources/'+source['id']+'/delete'
        with self.assertRaises(urllib.error.HTTPError) as error:self.req(endpoint)
        self.assertEqual(error.exception.code,405)
        with self.assertRaises(urllib.error.HTTPError) as error:self.req(endpoint,{'revision':1},{'Origin':'http://evil.example'})
        self.assertEqual(error.exception.code,403)
        with self.assertRaises(urllib.error.HTTPError) as error:self.req(endpoint,{})
        self.assertEqual(error.exception.code,400)
        self.assertEqual(self.req(endpoint,{'revision':1})['revision'],2)
        self.assertEqual(self.req(base+'/sources'),[])
        with self.assertRaises(urllib.error.HTTPError) as error:self.req('/api/jobs/'+job['id']+'/apply',{})
        self.assertEqual(error.exception.code,409)
        self.assertEqual(self.req(base)['project']['requirements'],[])

    def test_invalid_design_cannot_be_approved(self):
        p=self.req("/api/projects",{"sample":True});base="/api/projects/"+p["id"]
        p["project"]["design"]["devices"][1]["management_ip"]="192.0.2.11/24"
        self.req(base,{"revision":1,"project":p["project"]})
        with self.assertRaises(urllib.error.HTTPError) as error:self.req(base+"/approve",{"revision":2})
        self.assertEqual(error.exception.code,422)


if __name__=="__main__":unittest.main()
