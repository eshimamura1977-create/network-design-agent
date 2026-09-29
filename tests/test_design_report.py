import copy
import io
import json
import os
from pathlib import Path
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.request
from http.server import ThreadingHTTPServer
from unittest.mock import patch

from docx import Document

from agent.design_report import generate_design_report, local_draft, prepare_input
from agent.domain import Project
from agent.server import create_handler
from agent.store import Conflict, Store


class DesignReportTests(unittest.TestCase):
    def inputs(self):
        requirements = [dict(id='REQ-1', text='拠点LANのアドレスを維持する。', source='要件.txt',
            source_id='s1', quote='拠点LANのアドレスを維持する。', status='confirmed'),
            dict(id='REQ-2', text='RADIUSの冗長方式は未決。', source='要件.txt', source_id='s1',
                 quote='RADIUSの冗長方式は未決。', status='unresolved')]
        snapshot = dict(revision=3, approved_revision=None,
            project=Project(name='設計書検証', requirements=requirements).model_dump())
        sources = [dict(id='s1', name='要件.txt', body='L1: 拠点LANのアドレスを維持する。\nL2: RADIUSの冗長方式は未決。')]
        return snapshot, sources

    def test_local_word_preserves_requirements_without_mutation_or_ai(self):
        snapshot, sources = self.inputs()
        before = copy.deepcopy(snapshot)
        with tempfile.TemporaryDirectory() as tmp, patch('agent.design_report.design_chat_json') as chat:
            path = Path(tmp)/'report.docx'
            result = generate_design_report(snapshot, sources, path)
            chat.assert_not_called()
            text = '\n'.join(p.text for p in Document(path).paragraphs)
            for value in ['ネットワーク基本設計書', 'REQ-1 確認済み', 'REQ-2 未決', '未承認', 'WANと経路設計', 'RADIUSの冗長方式は未決。']:
                self.assertIn(value, text)
            self.assertEqual(result['chapters'], 10)
            self.assertEqual(result['mode'], 'local')
            self.assertEqual(snapshot, before)
            self.assertEqual(json.loads(path.with_suffix('.json').read_text(encoding='utf-8'))['sources'][0]['id'], 's1')

    def test_input_limits_and_duplicate_ids(self):
        snapshot, sources = self.inputs()
        with self.assertRaises(ValueError): prepare_input(dict(project=dict(requirements=[])), [])
        with self.assertRaisesRegex(ValueError, '5万字'): prepare_input(snapshot, [dict(body='a'*50001)])
        snapshot['project']['requirements'][1]['id'] = 'REQ-1'
        with self.assertRaisesRegex(ValueError, '重複'): prepare_input(snapshot, sources)

    def test_sources_without_reviewed_requirements_are_supported(self):
        snapshot, sources = self.inputs(); snapshot['project']['requirements'] = []
        prepare_input(snapshot, sources)
        draft = local_draft(snapshot, sources)
        self.assertTrue(any(s['evidence'] for s in draft['sections']))

    def test_ai_quote_ids_and_chapter_validation(self):
        snapshot, sources = self.inputs()
        original = local_draft(snapshot, sources)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/'valid.docx'
            result = generate_design_report(snapshot, sources, path, 'm365',
                lambda system, context: json.dumps(original, ensure_ascii=False))
            self.assertEqual(result['mode'], 'm365')
            for kind in ('quote', 'requirement', 'chapter', 'extra', 'question'):
                draft = copy.deepcopy(original)
                if kind == 'quote': draft['sections'][0]['evidence'] = [dict(source_id='s1', quote='存在しない原文')]
                if kind == 'requirement': draft['sections'][0]['requirement_ids'] = ['OTHER-PROJECT-ID']
                if kind == 'chapter': draft['sections'].pop()
                if kind == 'extra': draft['approved'] = True
                if kind == 'question': draft['sections'][0]['questions'] = ['x'*1001]
                invalid = Path(tmp)/(kind+'.docx')
                with self.subTest(kind=kind), self.assertRaises(ValueError):
                    generate_design_report(snapshot, sources, invalid, 'm365', lambda system, context: json.dumps(draft))
                self.assertFalse(invalid.exists())

    def test_no_silent_ai_fallback(self):
        snapshot, sources = self.inputs()
        with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ, {'LLM_MODEL':''}):
            for mode in ('ai', 'm365', 'unknown'):
                with self.subTest(mode=mode), self.assertRaises(ValueError):
                    generate_design_report(snapshot, sources, Path(tmp)/'invalid.docx', mode)

    def test_atomic_input_snapshot_and_revision_check(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(Path(tmp)); item = store.create(Project(name='snapshot').model_dump())
            sid = store.add_source(item['id'], 'source.txt', 'L1: VLAN分離')
            snap, sources = store.design_inputs(item['id'], 1)
            store.delete_source(item['id'], sid, 1)
            self.assertEqual(sources[0]['body'], 'L1: VLAN分離')
            self.assertEqual(snap['revision'], 1)
            with self.assertRaises(Conflict): store.design_inputs(item['id'], 1)
            with self.assertRaises(KeyError): store.design_inputs('0'*32, 1)

    def test_http_job_download_and_no_config_or_approval_change(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(Path(tmp)); hosts = set()
            server = ThreadingHTTPServer(('127.0.0.1', 0), create_handler(store, hosts))
            hosts.add(f'127.0.0.1:{server.server_port}')
            base = f'http://127.0.0.1:{server.server_port}'
            threading.Thread(target=server.serve_forever, daemon=True).start()
            def call(path, body=None, extra=None):
                headers={'Content-Type':'application/json','X-Workbench-Request':'1'}; headers.update(extra or {})
                req=urllib.request.Request(base+path, data=json.dumps(body).encode() if body is not None else None, headers=headers)
                with urllib.request.urlopen(req, timeout=10) as response:
                    return response.read(), response.headers
            try:
                item=store.create(Project(name='API生成検証').model_dump())
                endpoint='/api/projects/'+item['id']+'/design-document'
                store.add_source(item['id'], 'test.txt', 'L1: 社内VLANとゲストを分離する。')
                for body, headers, code in [({}, {}, 400), ({'revision':True}, {}, 400), ({'revision':9},{},409),
                    ({'revision':1,'mode':'m365'}, {}, 401), ({'revision':1}, {'Origin':'http://evil.example'},403)]:
                    with self.subTest(body=body), self.assertRaises(urllib.error.HTTPError) as error: call(endpoint, body, headers)
                    self.assertEqual(error.exception.code, code)
                with self.assertRaises(urllib.error.HTTPError) as error: call(endpoint)
                self.assertEqual(error.exception.code, 405)
                started=json.loads(call(endpoint, {'revision':1,'mode':'local'})[0])
                for _ in range(100):
                    job=store.job(started['job_id'])
                    if job['status'] not in ('queued','running'): break
                    time.sleep(.05)
                self.assertEqual(job['status'], 'completed', job)
                raw, headers=call(job['result']['download_url'])
                self.assertIn('wordprocessingml.document', headers['Content-Type'])
                self.assertIn('v1.docx', headers['Content-Disposition'])
                self.assertEqual(Document(io.BytesIO(raw)).paragraphs[0].text, 'ネットワーク基本設計書')
                self.assertEqual(store.get(item['id']), item)
                self.assertEqual(store.jobs(item['id'])[0]['kind'], 'design-document')
                self.assertFalse((Path(tmp)/'exports').exists())
            finally:
                server.shutdown(); server.server_close()


if __name__ == '__main__': unittest.main()
