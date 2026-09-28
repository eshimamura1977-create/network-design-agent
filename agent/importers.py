"""Local document text extraction. Never opens Office, executes macros or fetches links."""
from __future__ import annotations

import io
import posixpath
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

MAX_SOURCE_CHARS = 160000
TEXT_EXTENSIONS = ('.txt', '.md', '.csv', '.log', '.cfg')
SUPPORTED_EXTENSIONS = (*TEXT_EXTENSIONS, '.docx', '.xlsx', '.pptx', '.vsdx', '.vdx', '.pdf')
LEGACY_FORMATS = {'.doc': '.docx', '.xls': '.xlsx', '.ppt': '.pptx', '.vsd': '.vsdx'}


class Lines(list):
    def __init__(self):
        super().__init__()
        self.characters = 0

    def append(self, value):
        self.characters += len(value) + 1
        if self.characters > MAX_SOURCE_CHARS:
            raise ValueError('抽出文字数が16万字を超えています。資料を分割してください。')
        super().append(value)


def compact(value):
    return ' '.join(str(value).split())


def tag(element):
    return element.tag.rsplit('}', 1)[-1]


class NoDTD(ET.TreeBuilder):
    def doctype(self, name, pubid, system):
        raise ValueError('DTDを含むXMLには対応していません。標準形式で保存し直してください。')


def xml(data):
    return ET.fromstring(data, parser=ET.XMLParser(target=NoDTD()))


def check_archive(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 3000 or sum(x.file_size for x in entries) > 40 * 1024 * 1024:
            raise ValueError('Officeファイルの展開サイズが上限を超えています。必要な範囲を分割してください。')
        if len({x.filename for x in entries}) != len(entries):
            raise ValueError('Officeファイルに重複した部品があります。保存し直してください。')
        if any(x.flag_bits & 1 for x in entries):
            raise ValueError('暗号化されたOfficeファイルには対応していません。')


def relationships(archive, part):
    folder, filename = posixpath.split(part)
    relpart = posixpath.join(folder, '_rels', filename + '.rels')
    if relpart not in archive.namelist():
        return {}
    result = {}
    for rel in xml(archive.read(relpart)):
        if rel.get('TargetMode', '').lower() == 'external':
            continue
        target = rel.get('Target', '')
        if not target or '\\' in target or ':' in target:
            raise ValueError('Office部品の参照先が不正です。')
        resolved = posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(folder, target))
        if resolved.startswith('../') or resolved == '..':
            raise ValueError('Office部品の参照先がパッケージ外です。')
        result[rel.get('Id')] = (resolved, rel.get('Type', '').rsplit('/', 1)[-1])
    return result


def rel_id(node):
    return next((v for k, v in node.attrib.items() if k.endswith('}id')), None)


def paragraph_text(node):
    # Join runs without inserted spaces; preserve a line break as a space.
    return compact(''.join((x.text or '') if tag(x) == 't' else ' ' for x in node.iter() if tag(x) in ('t', 'br', 'tab')))


def ppt_text(archive, lines):
    part = 'ppt/presentation.xml'
    presentation = xml(archive.read(part))
    refs = relationships(archive, part)
    slide_ids = [x for x in presentation.iter() if tag(x) == 'sldId']
    if len(slide_ids) > 300:
        raise ValueError('PowerPointは300スライド以内にしてください。')
    for number, slide_id in enumerate(slide_ids, 1):
        target, kind = refs[rel_id(slide_id)]
        if kind != 'slide':
            raise ValueError('スライドの参照が不正です。')
        root = xml(archive.read(target))
        prefix = f'スライド{number}' + ('（非表示）' if root.get('show') == '0' else '')
        table_number = 0
        for shape in root.iter():
            if tag(shape) in ('sp', 'cxnSp'):
                identity = next((x.get('id', '?') for x in shape.iter() if tag(x) == 'cNvPr'), '?')
                paragraphs = [paragraph_text(x) for x in shape.iter() if tag(x) == 'p']
                content = ' / '.join(x for x in paragraphs if x)
                if content:
                    lines.append(f'{prefix} 図形{identity}: {content}')
            elif tag(shape) == 'tbl':
                table_number += 1
                rows = [x for x in shape if tag(x) == 'tr']
                for row_number, row in enumerate(rows, 1):
                    cells = [' / '.join(filter(None, (paragraph_text(p) for p in cell.iter() if tag(p) == 'p'))) for cell in row if tag(cell) == 'tc']
                    if any(cells):
                        lines.append(f'{prefix} 表{table_number} 行{row_number}: ' + ' | '.join(cells))
        for note_part, note_kind in relationships(archive, target).values():
            if note_kind != 'notesSlide':
                continue
            notes = xml(archive.read(note_part))
            for shape in (x for x in notes.iter() if tag(x) == 'sp'):
                placeholders = [x.get('type') for x in shape.iter() if tag(x) == 'ph']
                if placeholders and 'body' not in placeholders:
                    continue
                content = ' / '.join(filter(None, (paragraph_text(x) for x in shape.iter() if tag(x) == 'p')))
                if content:
                    lines.append(f'{prefix} ノート: {content}')


def visio_page(root, prefix, lines):
    for shape in (x for x in root.iter() if tag(x) == 'Shape'):
        sid = shape.get('ID', '?')
        name = shape.get('NameU', shape.get('Name', ''))
        label = f'{prefix} 図形{sid}'
        texts = [compact(''.join(x.itertext())) for x in shape if tag(x) == 'Text']
        content = ' / '.join(x for x in texts if x)
        if content:
            lines.append(f'{label}: {content}')
        if name:
            lines.append(f'{label} 名前: {name}')
        for section in (x for x in shape if tag(x) == 'Section' and x.get('N') == 'Property'):
            for row in section:
                if tag(row) != 'Row' or row.get('Del') == '1':
                    continue
                cells = {c.get('N'): c for c in row if tag(c) == 'Cell'}
                value = cells.get('Value')
                if value is not None:
                    caption = cells.get('Label')
                    key = caption.get('V', '') if caption is not None else row.get('N', row.get('IX', '?'))
                    lines.append(f'{label} 属性{compact(key)}: {compact(value.get("V", ""))}')
        # Legacy XML VDX stores shape data in Prop elements instead of Section.
        for prop in (x for x in shape if tag(x) == 'Prop'):
            values = {tag(x): compact(''.join(x.itertext())) for x in prop}
            if values.get('Value'):
                lines.append(f'{label} 属性{values.get("Label", prop.get("NameU", "?"))}: {values["Value"]}')
        if shape.get('Master') is not None or shape.get('MasterShape') is not None:
            lines.append(f'{label} 継承参照: Master={shape.get("Master", "親から継承")} MasterShape={shape.get("MasterShape", "-")}（継承文字・属性は未展開）')
    for i, connect in enumerate((x for x in root.iter() if tag(x) == 'Connect'), 1):
        lines.append(f'{prefix} 接続{i}: 図形{connect.get("FromSheet", "?")}.{connect.get("FromCell", "?")} → 図形{connect.get("ToSheet", "?")}.{connect.get("ToCell", "?")}（Visio登録情報。物理ポートの推定なし）')


def visio_text(data, lines, zipped=True):
    if not zipped:
        pages = [x for x in xml(data).iter() if tag(x) == 'Page']
        if len(pages) > 300:
            raise ValueError('Visioは300ページ以内にしてください。')
        for i, page in enumerate(pages, 1):
            visio_page(page, f'ページ{i}({page.get("Name", page.get("NameU", "無題"))})', lines)
        return
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        part = 'visio/pages/pages.xml'
        pages = [x for x in xml(archive.read(part)) if tag(x) == 'Page']
        if len(pages) > 300:
            raise ValueError('Visioは300ページ以内にしてください。')
        refs = relationships(archive, part)
        for i, page in enumerate(pages, 1):
            rid = next((rel_id(x) for x in page if tag(x) == 'Rel'), None)
            target, kind = refs[rid]
            if kind != 'page':
                raise ValueError('Visioページの参照が不正です。')
            prefix = f'ページ{i}({page.get("Name", page.get("NameU", "無題"))})'
            if page.get('Background') == '1':
                prefix += '（背景）'
            visio_page(xml(archive.read(target)), prefix, lines)


def excel_text(data, lines):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    book = load_workbook(io.BytesIO(data), data_only=False, read_only=True, keep_links=False)
    cached = None
    try:
        cached = load_workbook(io.BytesIO(data), data_only=True, read_only=True, keep_links=False)
        budget = 0
        for sheet in book:
            rows, cols = sheet.max_row or 0, sheet.max_column or 0
            budget += rows * cols
            if rows > 20000 or cols > 256 or budget > 400000:
                raise ValueError('Excelは各シート2万行・256列、全シート合計40万セル以内にしてください。')
            prefix = sheet.title + ('（非表示）' if sheet.sheet_state != 'visible' else '')
            for rownum, (row, values) in enumerate(zip(sheet.iter_rows(), cached[sheet.title].iter_rows()), 1):
                parts = []
                for colnum, (cell, saved) in enumerate(zip(row, values), 1):
                    if cell.value is None:
                        continue
                    value = compact(cell.value)
                    if cell.data_type == 'f':
                        value = f'{value} [保存済み値={compact(saved.value)}]' if saved.value is not None else f'{value} [数式・保存済み値なし／未再計算]'
                    parts.append(f'{get_column_letter(colnum)}{rownum}={value}')
                if parts:
                    lines.append(f'{prefix}: ' + ' | '.join(parts))
    finally:
        book.close()
        if cached:
            cached.close()


def extract_file(name: str, data: bytes) -> str:
    if len(data) > 8 * 1024 * 1024:
        raise ValueError('ファイルは8MB以内にしてください')
    ext = Path(name).suffix.lower()
    if ext in LEGACY_FORMATS:
        raise ValueError(f'旧形式{ext}は直接取込できません。元のアプリで{LEGACY_FORMATS[ext]}形式に「名前を付けて保存」して取り込んでください。拡張子の変更だけでは変換できません。')
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError('対応形式はDOCX / XLSX / PPTX / VSDX / VDX / TXT / MD / CSV / CFG / LOG / 文字を含むPDFです。')
    lines = Lines()
    try:
        if ext in ('.docx', '.xlsx', '.pptx', '.vsdx'):
            check_archive(data)
        if ext in TEXT_EXTENSIONS:
            if data.startswith((b'\xff\xfe', b'\xfe\xff')):
                value = data.decode('utf-16')
            else:
                try:
                    value = data.decode('utf-8-sig')
                except UnicodeDecodeError:
                    value = data.decode('cp932')
            for i, line in enumerate(value.splitlines(), 1):
                if line.strip():
                    lines.append(f'L{i}: {line}')
        elif ext == '.docx':
            from docx import Document
            doc = Document(io.BytesIO(data))
            for i, paragraph in enumerate(doc.paragraphs, 1):
                if paragraph.text.strip():
                    lines.append(f'段落{i}: {paragraph.text}')
            for i, table in enumerate(doc.tables, 1):
                for j, row in enumerate(table.rows, 1):
                    lines.append(f'表{i} 行{j}: ' + ' | '.join(compact(c.text) for c in row.cells))
        elif ext == '.xlsx':
            excel_text(data, lines)
        elif ext == '.pptx':
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                ppt_text(archive, lines)
        elif ext in ('.vsdx', '.vdx'):
            visio_text(data, lines, zipped=ext == '.vsdx')
        elif ext == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted or len(reader.pages) > 150:
                raise ValueError('暗号化PDFまたは150ページ超のPDFには対応していません')
            for i, page in enumerate(reader.pages, 1):
                for j, line in enumerate((page.extract_text() or '').splitlines(), 1):
                    if line.strip():
                        lines.append(f'P{i} L{j}: {line}')
    except (zipfile.BadZipFile, KeyError, ET.ParseError, UnicodeError, EOFError):
        raise ValueError('資料を読み取れません。破損・暗号化・拡張子と実形式の不一致を確認し、元のアプリで保存し直してください。') from None
    result = '\n'.join(lines)
    if not result.strip():
        raise ValueError('文字を抽出できません。画像のみの資料はOCR済みテキストを使用してください。')
    return result
