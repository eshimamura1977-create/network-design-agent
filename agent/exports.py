from __future__ import annotations

import hashlib
import html
import json
import math
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from .domain import config_text, validate_project


def topology_svg(design, test=False):
    devices = design["devices"]
    width = 1050
    height = max(360, 160 + math.ceil(max(0, len(devices) - 1) / 3) * 160)
    positions = {}
    for i, dev in enumerate(devices):
        positions[dev["hostname"]] = (525, 82) if i == 0 else (175 + ((i - 1) % 3) * 350, 250 + ((i - 1) // 3) * 160)
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height}" role="img" aria-label="ネットワーク構成図"><rect width="100%" height="100%" fill="#f5f8fc"/>']
    for link in design["links"]:
        if link["a_device"] not in positions or link["b_device"] not in positions:
            continue
        x1, y1 = positions[link["a_device"]]
        x2, y2 = positions[link["b_device"]]
        label = f"{link['a_port'].replace('GigabitEthernet', 'Gi')} ↔ {link['b_port'].replace('GigabitEthernet', 'Gi')}"
        out.append(f'<path d="M{x1},{y1+43} L{x2},{y2-43}" stroke="#829ab5" stroke-width="2" fill="none"/>')
        out.append(f'<text x="{(x1+x2)/2}" y="{(y1+y2)/2}" font-family="sans-serif" font-size="12" text-anchor="middle" fill="#48627e">{html.escape(label)}</text>')
    for dev in devices:
        x, y = positions[dev["hostname"]]
        out += [f'<rect x="{x-125}" y="{y-43}" width="250" height="86" rx="10" fill="white" stroke="#a8bdcf"/>', f'<rect x="{x-125}" y="{y-43}" width="5" height="86" rx="2" fill="#1769e0"/>']
        for offset, text, size, color in [(-15, dev["hostname"], 16, "#17334c"), (6, dev["model"], 12, "#5a7186"), (26, dev["management_ip"], 13, "#5a7186")]:
            out.append(f'<text x="{x}" y="{y+offset}" text-anchor="middle" font-family="sans-serif" font-size="{size}" fill="{color}">{html.escape(text)}</text>')
    note = "試験構成の下書き：本番と同じ接続。測定端末・負荷装置・障害注入点は未定義。" if test else "L2接続図。ゲートウェイ・端末・外部装置は本図の対象外。"
    out.append(f'<text x="20" y="{height-12}" font-family="sans-serif" font-size="12" fill="#586a7b">{note}</text></svg>')
    return "".join(out)


def make_vsdx(path: Path, design: dict, test=False):
    """Minimal editable VSDX. Geometry is deterministic; Visio desktop QA remains required."""
    ns = "http://schemas.microsoft.com/office/visio/2012/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ET.register_namespace("", ns)
    ET.register_namespace("r", rel)
    def node(tag, attrs=None):
        return ET.Element("{" + ns + "}" + tag, attrs or {})
    def child(parent, tag, attrs=None):
        item = node(tag, attrs)
        parent.append(item)
        return item
    def cell(parent, n, value, **extra):
        return child(parent, "Cell", dict(N=n, V=str(value), **extra))
    def xml(e):
        return ET.tostring(e, encoding="utf-8", xml_declaration=True)
    height = max(5.0, 2.2 + math.ceil(max(0, len(design["devices"]) - 1) / 3) * 1.7)
    page = node("PageContents")
    shapes = child(page, "Shapes")
    coords, ids = {}, {}
    for i, dev in enumerate(design["devices"], 1):
        x, y = (5.5, height - 0.8) if i == 1 else (1.8 + ((i - 2) % 3) * 3.7, height - 2.6 - ((i - 2) // 3) * 1.7)
        coords[dev["hostname"]] = (x, y)
        ids[dev["hostname"]] = i
        shape = child(shapes, "Shape", dict(ID=str(i), NameU=f"Device.{i}", Type="Shape", LineStyle="0", FillStyle="0", TextStyle="0"))
        for k, v in dict(PinX=x, PinY=y, Width=2.6, Height=0.9, LocPinX=1.3, LocPinY=0.45, LineColor="#91A9BD", LinePattern=1, FillForegnd="#FFFFFF", FillPattern=1).items():
            cell(shape, k, v)
        geo = child(shape, "Section", dict(N="Geometry", IX="0"))
        for j, (typ, a, b) in enumerate([("MoveTo", 0, 0), ("LineTo", 2.6, 0), ("LineTo", 2.6, .9), ("LineTo", 0, .9), ("LineTo", 0, 0)], 1):
            row = child(geo, "Row", dict(T=typ, IX=str(j)))
            cell(row, "X", a); cell(row, "Y", b)
        child(shape, "Text").text = f"{dev['hostname']}\n{dev['model']}\n{dev['management_ip']}"
    connects = child(page, "Connects")
    for i, link in enumerate(design["links"], len(design["devices"]) + 1):
        if link["a_device"] not in coords or link["b_device"] not in coords:
            continue
        x1, y1 = coords[link["a_device"]]; x2, y2 = coords[link["b_device"]]
        y1 -= .45; y2 += .45
        dx, dy = x2-x1, y2-y1
        length = math.hypot(dx, dy)
        s = child(shapes, "Shape", dict(ID=str(i), NameU=f"Link.{i}", Type="Shape", LineStyle="0", FillStyle="0", TextStyle="0"))
        for k, v in dict(BeginX=x1, BeginY=y1, EndX=x2, EndY=y2, PinX=(x1+x2)/2, PinY=(y1+y2)/2, Width=length, Height=0, LocPinX=length/2, LocPinY=0, Angle=math.atan2(dy,dx), OneD=1, LinePattern=1, LineColor="#647D94", FillPattern=0).items():
            cell(s,k,v)
        geo = child(s, "Section", dict(N="Geometry", IX="0"))
        for j, (typ, x) in enumerate([("MoveTo",0), ("LineTo",length)],1):
            row=child(geo,"Row",dict(T=typ,IX=str(j))); cell(row,"X",x); cell(row,"Y",0)
        child(s,"Text").text = link["a_port"].replace("GigabitEthernet","Gi") + " - " + link["b_port"].replace("GigabitEthernet","Gi")
        for cell_name,host,part in [("BeginX",link["a_device"],"9"),("EndX",link["b_device"],"12")]:
            child(connects,"Connect",dict(FromSheet=str(i),FromCell=cell_name,FromPart=part,ToSheet=str(ids[host]),ToCell="PinX",ToPart="3"))
    doc=node("VisioDocument")
    styles=child(doc,"StyleSheets")
    style=child(styles,"StyleSheet",dict(ID="0",NameU="No Style"))
    for k,v in dict(LinePattern=1,FillPattern=1,FillForegnd="#FFFFFF").items(): cell(style,k,v)
    pages=node("Pages")
    p=child(pages,"Page",dict(ID="0",NameU="Test topology draft" if test else "L2 topology",Name="試験構成下書き" if test else "L2構成図"))
    sheet=child(p,"PageSheet")
    cell(sheet,"PageWidth",11); cell(sheet,"PageHeight",height)
    child(p,"Rel",{"{"+rel+"}id":"rId1"})
    relation_ns="http://schemas.openxmlformats.org/package/2006/relationships"
    def relationships(type_name,target):
        return f'<?xml version="1.0" encoding="utf-8"?><Relationships xmlns="{relation_ns}"><Relationship Id="rId1" Type="http://schemas.microsoft.com/visio/2010/relationships/{type_name}" Target="{target}"/></Relationships>'
    types='''<?xml version="1.0" encoding="utf-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/visio/document.xml" ContentType="application/vnd.ms-visio.drawing.main+xml"/><Override PartName="/visio/pages/pages.xml" ContentType="application/vnd.ms-visio.pages+xml"/><Override PartName="/visio/pages/page1.xml" ContentType="application/vnd.ms-visio.page+xml"/></Types>'''
    with zipfile.ZipFile(path,"w",zipfile.ZIP_DEFLATED) as z:
        for name,data in {"[Content_Types].xml":types,"_rels/.rels":relationships("document","visio/document.xml"),"visio/document.xml":xml(doc),"visio/_rels/document.xml.rels":relationships("pages","pages/pages.xml"),"visio/pages/pages.xml":xml(pages),"visio/pages/_rels/pages.xml.rels":relationships("page","page1.xml"),"visio/pages/page1.xml":xml(page)}.items(): z.writestr(name,data)


def test_rows(project):
    rows=[]
    for r in project["requirements"]:
        rows.append([f"T-{len(rows)+1:03d}",r["id"],"要件受入",r["text"],"確認方法・対象・閾値をレビューで具体化する",r["text"],"未実施","要具体化"])
    for dev in project["design"]["devices"]:
        rows.append([f"T-{len(rows)+1:03d}","設計データ",dev["hostname"],"管理IP設定", "show ip interface brief と台帳を照合", dev["management_ip"],"未実施","手順案"])
        for p in dev["ports"]:
            rows.append([f"T-{len(rows)+1:03d}","設計データ",dev["hostname"]+":"+p["name"],"L2ポート設定",f"show interfaces {p['name']} switchport と台帳を照合",f"mode={p['mode']}, VLAN={','.join(map(str,p['vlans']))}"+(f", native={p['native_vlan']}" if p['mode']=="trunk" else ""),"未実施","手順案"])
    return rows


def write_xlsx(path, sheets):
    # Application runtime exporter: uses public, installable dependencies.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    wb=Workbook(); wb.remove(wb.active)
    for name,headers,rows,widths in sheets:
        safe=re.sub(r'[\\/*?:\[\]]','_',name)[:31]
        ws=wb.create_sheet(safe)
        ws.append(headers)
        for values in rows:
            ws.append(values)
        for row in ws:
            for c in row:
                if isinstance(c.value,str):
                    c.data_type="s"  # Imported values must never become spreadsheet formulas.
                c.font=Font(name="Yu Gothic",size=10,color="203A50")
                c.alignment=Alignment(vertical="top",wrap_text=True)
                c.border=Border(bottom=Side(style="hair",color="DCE5ED"))
                if c.row>1 and c.row%2==0: c.fill=PatternFill("solid",fgColor="F1F6FA")
        for c in ws[1]:
            c.fill=PatternFill("solid",fgColor="153C59"); c.font=Font(name="Yu Gothic",size=10,bold=True,color="FFFFFF")
        ws.row_dimensions[1].height=30
        from openpyxl.utils import get_column_letter
        for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
        for idx,values in enumerate(rows,2):
            lines=max([1]+[math.ceil(len(str(v or ""))/max(1,widths[i]//2)) for i,v in enumerate(values)])
            ws.row_dimensions[idx].height=max(30,min(300,lines*16+8))
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        ws.sheet_view.showGridLines=False
        ws.print_title_rows="1:1"; ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation="landscape"; ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
        if "結果" in headers and rows:
            col=get_column_letter(headers.index("結果")+1)
            dv=DataValidation(type="list",formula1='"未実施,合格,不合格,保留"'); ws.add_data_validation(dv); dv.add(f"{col}2:{col}{len(rows)+1}")
    wb.save(path)


def write_doc(path,title,snapshot,sections):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.65); sec.left_margin=sec.right_margin=Inches(.7)
    for style in ["Normal","Title","Heading 1","Heading 2"]:
        s=doc.styles[style]; s.font.name="Yu Gothic"; s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Yu Gothic")
    doc.styles["Normal"].font.size=Pt(10)
    doc.styles["Normal"].paragraph_format.space_after=Pt(6)
    doc.styles["Title"].font.size=Pt(22)
    doc.add_paragraph(title,"Title")
    doc.add_paragraph(snapshot["project"]["name"])
    doc.add_paragraph(f"設計データ v{snapshot['revision']}　作成状態 下書き")
    doc.add_paragraph("登録済みの要件と設計データから作成したレビュー用資料です。未決事項を確認し、実機検証および担当者のレビュー後に確定してください。")
    for heading,paragraphs,headers,rows in sections:
        doc.add_heading(heading,level=1)
        for text in paragraphs: doc.add_paragraph(str(text))
        if headers:
            table=doc.add_table(rows=1,cols=len(headers)); table.style="Table Grid"
            for c,t in zip(table.rows[0].cells,headers): c.text=str(t)
            for c in table.rows[0].cells:
                shade=OxmlElement("w:shd"); shade.set(qn("w:fill"),"E7EEF4"); c._tc.get_or_add_tcPr().append(shade)
            repeat=OxmlElement("w:tblHeader"); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for values in rows:
                for c,t in zip(table.add_row().cells,values): c.text=str(t)
            for row in table.rows:
                for c in row.cells:
                    for p in c.paragraphs:
                        p.paragraph_format.space_after=Pt(4)
                        for run in p.runs: run.font.size=Pt(9)
            doc.add_paragraph()
    sec.footer.paragraphs[0].text=f"Network Design Workbench  •  v{snapshot['revision']}  •  DRAFT"
    doc.save(path)


def export_bundle(snapshot, folder: Path):
    folder.mkdir(parents=True,exist_ok=False)
    p=snapshot["project"]; d=p["design"]; rev=snapshot["revision"]
    issues=validate_project(p)
    approved=snapshot.get("approved_revision")==rev
    eligible=approved and not any(x["level"] in ("error","warning") for x in issues)
    (folder/"design-data.json").write_text(json.dumps(snapshot,ensure_ascii=False,indent=2),encoding="utf-8")
    (folder/"validation.json").write_text(json.dumps(issues,ensure_ascii=False,indent=2),encoding="utf-8")
    reqsections=[("要件と根拠",[],None,[])]
    statuses={"candidate":"候補","confirmed":"確認済み","unresolved":"未決"}
    for r in p["requirements"]:
        reqsections.append((r["id"]+" "+r["category"],[r["text"],"状態 "+statuses[r["status"]],"出典 "+r["source"],"原文 "+(r["quote"] or "引用未登録")],None,[]))
    write_doc(folder/"01_要件定義書.docx","ネットワーク要件定義書",snapshot,reqsections)
    device_rows=[[x["hostname"],x["model"],x["management_ip"]] for x in d["devices"]]
    vlan_rows=[[v["id"],v["name"],v["subnet"],v["purpose"]] for v in d["vlans"]]
    sections=[("設計範囲と未決事項",[d["notes"] or "設計方針は未入力です。", "本試作のConfig生成範囲はC9200LのL2 VLAN、管理SVI、アクセスポート、トランクです。"],None,[]),("機器構成",[],["ホスト名","機種","管理IP"],device_rows),("VLANとアドレス",[],["VLAN","名称","サブネット","用途"],vlan_rows),("静的検査結果",[x["message"] for x in issues],None,[])]
    for dev in d["devices"]:
        sections.append((dev["hostname"]+" の接続設定",[],["インターフェース","モード","VLAN"],[[a["name"],a["mode"],",".join(map(str,a["vlans"]))] for a in dev["ports"]]))
    write_doc(folder/"02_ネットワーク設計書.docx","ネットワーク設計書",snapshot,sections)
    overview=[["案件",p["name"]],["設計版",rev],["設計承認","承認済み" if approved else "未承認"],["用途","生成時点のスナップショット。Excel変更はアプリへ自動反映されません。"]]
    sheets=[("概要",["項目","値"],overview,[24,85]),("アドレス台帳",["ホスト名","VRF","管理IP","管理VLAN","ゲートウェイ"],[[x["hostname"],x["vrf"],x["management_ip"],x["management_vlan"],x["gateway"]] for x in d["devices"]],[25,18,26,18,22]),("VLAN台帳",["VLAN","名称","サブネット","VRF","用途"],[[v["id"],v["name"],v["subnet"],v["vrf"],v["purpose"]] for v in d["vlans"]],[12,25,26,18,40]),("接続台帳",["接続元","ポート","接続先","ポート"],[[l[k] for k in ["a_device","a_port","b_device","b_port"]] for l in d["links"]],[24,30,24,30])]
    for index,dev in enumerate(d["devices"],1):
        sheets.append((f"{index:02d}_{dev['hostname']}",["ポート","モード","VLAN","Native VLAN","説明"],[[x["name"],x["mode"],",".join(map(str,x["vlans"])),x["native_vlan"] if x["mode"]=="trunk" else "対象外",x["description"]] for x in dev["ports"]],[30,16,25,18,40]))
    for index,model in enumerate(sorted(set(x["model"] for x in d["devices"])),1):
        sheets.append((f"機種{index}_{model}",["ホスト名","機種","OS版","役割"],[[x["hostname"],x["model"],x["os_version"],x["role"]] for x in d["devices"] if x["model"]==model],[25,25,25,25]))
    write_xlsx(folder/"03_パラメータシート.xlsx",sheets)
    for prefix,test in [("04_構成図",False),("07_試験構成図",True)]:
        (folder/(prefix+".svg")).write_text(topology_svg(d,test),encoding="utf-8")
        make_vsdx(folder/(prefix+"_試作.vsdx"),d,test)
    if eligible:
        cfg=folder/"05_Config"; cfg.mkdir()
        for dev in d["devices"]:
            (cfg/(dev["hostname"]+".cfg")).write_text(f"! Source design revision: {rev}\n"+config_text(dev,d),encoding="utf-8")
    else:
        (folder/"05_Config_未出力.txt").write_text("設計の承認とエラー・警告の解消後に再生成してください。Configは出力していません。",encoding="utf-8")
    write_doc(folder/"06_試験設計書.docx","ネットワーク試験設計書",snapshot,[("目的と試験範囲",["要件および登録済みL2設定が期待値に一致することを確認します。", "試験項目表の『要具体化』行は実施前に手順・対象・判定閾値を確定してください。"],None,[]),("試験環境",["構成図は登録された接続を転記しています。測定端末・負荷装置・障害注入点は未定義です。", "機器の正確なSKU、OS版、ライセンス、試験用アドレスを別途確定します。"],None,[]),("実施と判定",["設定確認、疎通、分離、異常時の順で試験を設計します。業務影響のある操作は検証環境で実施します。", "各項目に実施日時、実測結果、ログ、実施者、判定を記録します。未実施を合格として扱いません。"],None,[])])
    write_xlsx(folder/"08_試験項目表.xlsx",[("試験項目",["ID","要件ID","対象","試験観点","手順案","期待値","結果","具体化状態"],test_rows(p),[15,20,30,30,60,55,16,18])])
    write_doc(folder/"09_移行設計書.docx","ネットワーク移行設計書",snapshot,[("移行方式",["現行構成との差分、業務停止許容時間、移行順序は未確定です。入力資料と関係者の確認に基づき具体化してください。", "この資料は移行計画の下書きであり、そのまま実行できる作業手順ではありません。"],None,[]),("事前準備",["現行Config・配線・接続先の採取、バックアップの復元確認、連絡先・承認者の確定を行います。"],None,[]),("切替と切戻し",["切替順序、疎通確認、業務確認、切戻し判断時刻を決定します。", "切戻しの開始条件、責任者、復旧手順、復旧所要時間は未確定です。"],None,[])])
    work=[ [1,"事前準備","現行設定・配線・接続先を記録","バックアップと現地情報を照合","未定","未実施"], [2,"事前確認","承認版・対象装置・作業窓を確認","版番号と現物の一致","未定","未実施"], [3,"切替","確定した手順に従って実施","具体的コマンド・順序は未確定","未定","未実施"], [4,"確認","試験項目表を実施","受入条件を満たす","未定","未実施"], [5,"切戻し判断","判断時刻・条件に照らして判断","条件と責任者は未確定","未定","未実施"] ]
    write_xlsx(folder/"10_工事資料.xlsx",[("作業手順下書き",["順序","工程","作業","確認事項","担当","結果"],work,[12,20,55,60,18,16])])
    write_doc(folder/"11_運用引継ぎ資料.docx","ネットワーク運用引継ぎ資料",snapshot,[("対象と管理情報",["管理アドレスと装置一覧は同じ設計版のパラメータシートを参照してください。認証情報はこの資料に保存せず、承認された秘密情報管理先を使用します。"],None,[]),("運用上の未決事項",["監視対象、閾値、通知先、バックアップ周期、保管先、保守窓口は未確定です。", "障害受付、一次切分け、エスカレーション、復旧確認の担当を確定してください。"],None,[]),("確認コマンドの例",["show version、show interfaces status、show vlan brief、show interfaces trunk、show spanning-treeを参考に、実機OSで利用可能な手順を確定してください。"],None,[])])
    warnings=["全成果物は下書きです。ConfigはL2の一部設定のみで、実機投入・接続は行いません。", "Visio VSDXは基本図形での試作出力です。Visio実機での表示・コネクタ追従は未検証。SVGを併記しています。", "試験構成は設計構成を転記した下書きです。測定装置・故障注入点は未設計。", "移行・工事・引継ぎは未決事項を明示したひな型です。現行資料からの自動設計は未実装。"]
    (folder/"README.txt").write_text("\n".join(warnings)+f"\n設計版: {rev}\nConfig出力: {eligible}",encoding="utf-8")
    files=[]
    for file in sorted(folder.rglob("*")):
        if file.is_file():
            files.append(dict(name=file.relative_to(folder).as_posix(),size=file.stat().st_size,sha256=hashlib.sha256(file.read_bytes()).hexdigest()))
    manifest=dict(revision=rev,config_included=eligible,warnings=warnings,files=files)
    (folder/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
    archive=folder.with_suffix(".zip")
    with zipfile.ZipFile(archive,"w",zipfile.ZIP_DEFLATED) as z:
        for file in sorted(folder.rglob("*")):
            if file.is_file(): z.write(file,file.relative_to(folder).as_posix())
    return manifest
